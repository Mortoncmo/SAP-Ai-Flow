import json
import re
from hashlib import sha256
from collections.abc import Iterator
from datetime import UTC, datetime, timedelta
from io import BytesIO
from pathlib import Path
from threading import Event, Thread
from time import sleep
from urllib.parse import unquote

import pytest
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.shared import Inches
from fastapi.testclient import TestClient
from sqlalchemy import event, func, select
from sqlalchemy.exc import OperationalError

from app.agent.base import ProviderResult
from app.agent.result_cache import ModelResultCache
from app.api.routes.flowcharts import get_provider
from app.api.routes.projects import get_model_result_cache, get_repository
from app.core.config import Settings, get_settings
from app.core.errors import FlowchartError, ProviderError
from app.db.database import Database
from app.db.models import (
    ChangeLogRecord,
    ExportJobRecord,
    GapDecisionRecord,
    ProcessRecord,
    ProcessRevisionRecord,
    ProjectAuditRecord,
    ProjectMemberRecord,
    ProjectRecord,
)
from app.db.repository import BlueprintRepository
from app.documents import export_service
from app.documents.artifact_store import FilesystemArtifactStore
from app.knowledge.service import get_knowledge_service
from app.main import app
from app.models.graph import SapContext
from app.models.patch import LLMPatch
from app.workers.export_worker import ExportJobWorker


class TrackingExternalProvider:
    external = True

    def __init__(self) -> None:
        self.calls = 0
        self.evidence = []

    async def generate_patch(self, graph, instruction, locale, evidence=None):
        del graph, instruction, locale
        self.calls += 1
        self.evidence = evidence or []
        return ProviderResult(
            patch=LLMPatch.model_validate(
                {
                    "change_summary": (
                        "供应商：华东机密供应商，联系人：张三，"
                        "邮箱 secret@example.com，金额 100万元"
                    ),
                    "operations": [
                        {
                            "op": "add_node",
                            "ref": f"external_node_{self.calls}",
                            "node": {
                                "type": "task",
                                "label": "供应商：华东机密供应商",
                            },
                        }
                    ],
                }
            ),
            provider="fake-external",
            model="fake-external-v1",
        )


class UngroundedProvider:
    external = False

    async def generate_patch(self, graph, instruction, locale, evidence=None):
        del graph, instruction, locale, evidence
        return ProviderResult(
            patch=LLMPatch.model_validate(
                {
                    "change_summary": "增加未经证据支持的交易码。",
                    "operations": [
                        {
                            "op": "add_node",
                            "ref": "unverified",
                            "node": {
                                "type": "task",
                                "label": "虚构步骤",
                                "sap": {
                                    "tcodes": [
                                        {
                                            "code": "ZFAKE",
                                            "status": "verified",
                                            "evidence_ref": "invented#zfake",
                                        }
                                    ]
                                },
                            },
                        }
                    ],
                }
            ),
            provider="ungrounded",
            model="test",
        )


class ExplodingKnowledgeService:
    def search(self, *_args, **_kwargs):
        raise FlowchartError(
            "KNOWLEDGE_UNAVAILABLE",
            "知识库暂时不可用，请稍后重试。",
            status_code=503,
        )


class ExplodingProvider:
    external = False

    async def generate_patch(self, *_args, **_kwargs):
        raise ProviderError("PROVIDER_UNAVAILABLE", "模型服务暂时不可用，请稍后重试。")


@pytest.fixture
def persistence_client(tmp_path) -> Iterator[tuple[TestClient, Database]]:
    database = Database(f"sqlite:///{(tmp_path / 'blueprint.db').as_posix()}")

    def override_repository() -> Iterator[BlueprintRepository]:
        session = database.session_factory()
        try:
            yield BlueprintRepository(session)
        finally:
            session.close()

    app.dependency_overrides[get_repository] = override_repository
    try:
        with TestClient(app) as client:
            yield client, database
    finally:
        app.dependency_overrides.pop(get_repository, None)


def _create_process(client: TestClient) -> str:
    project_response = client.post(
        "/api/v1/projects",
        json={"name": "采购蓝图验收项目", "customer_name": "示例客户"},
    )
    assert project_response.status_code == 201, project_response.text
    project_id = project_response.json()["id"]

    process_response = client.post(
        f"/api/v1/projects/{project_id}/processes",
        json={"name": "直接物料 P2P", "module": "MM", "process_scope": "P2P"},
    )
    assert process_response.status_code == 201, process_response.text
    assert process_response.json()["current_revision"] == 0
    return process_response.json()["id"]


def test_drawio_source_save_persists_projection_hash_and_audit(persistence_client):
    client, database = persistence_client
    process_id = _create_process(client)
    current = client.get(f"/api/v1/processes/{process_id}").json()
    graph = {**current["graph"], "version": 1}
    xml = (
        '<mxfile host="SAP AI Flow"><diagram name="直接物料 P2P">'
        '<mxGraphModel><root><mxCell id="0"/><mxCell id="1" parent="0"/>'
        '</root></mxGraphModel></diagram></mxfile>'
    )
    digest = sha256(xml.encode("utf-8")).hexdigest()

    saved = client.post(
        f"/api/v1/processes/{process_id}/drawio",
        json={
            "request_id": "drawio-save-1",
            "base_revision": 0,
            "base_sha256": None,
            "xml": xml,
            "xml_sha256": digest,
            "graph": graph,
            "summary": "Draw.io 编辑器保存",
        },
    )
    assert saved.status_code == 200, saved.text
    assert saved.json()["revision_no"] == 1
    assert saved.json()["xml_sha256"] == digest

    loaded = client.get(f"/api/v1/processes/{process_id}/revisions/1/drawio")
    assert loaded.status_code == 200, loaded.text
    assert loaded.json()["xml"] == xml
    assert loaded.json()["graph"] == graph

    with database.session_factory() as session:
        revision = session.scalar(
            select(ProcessRevisionRecord).where(
                ProcessRevisionRecord.process_id == process_id,
                ProcessRevisionRecord.revision_no == 1,
            )
        )
        audit = session.scalar(
            select(ChangeLogRecord).where(
                ChangeLogRecord.process_id == process_id,
                ChangeLogRecord.result_revision == 1,
            )
        )
        assert revision is not None and revision.drawio_sha256 == digest
        assert audit is not None and audit.normalized_patch["kind"] == "drawio_save"


def test_drawio_source_save_rejects_hash_mismatch_and_stale_revision(persistence_client):
    client, _database = persistence_client
    process_id = _create_process(client)
    graph = client.get(f"/api/v1/processes/{process_id}").json()["graph"]
    graph["version"] = 1
    xml = '<mxfile><diagram><mxGraphModel><root/></mxGraphModel></diagram></mxfile>'
    payload = {
        "request_id": "drawio-save-conflict",
        "base_revision": 0,
        "base_sha256": None,
        "xml": xml,
        "xml_sha256": "0" * 64,
        "graph": graph,
    }

    mismatch = client.post(f"/api/v1/processes/{process_id}/drawio", json=payload)
    assert mismatch.status_code == 422
    assert mismatch.json()["error"]["code"] == "DRAWIO_HASH_MISMATCH"

    payload["xml_sha256"] = sha256(xml.encode("utf-8")).hexdigest()
    first = client.post(f"/api/v1/processes/{process_id}/drawio", json=payload)
    assert first.status_code == 200, first.text
    stale = client.post(f"/api/v1/processes/{process_id}/drawio", json=payload)
    assert stale.status_code == 409
    assert stale.json()["error"]["code"] == "REVISION_CONFLICT"


def test_modify_preview_returns_llm_patch_without_persisting_revision(persistence_client):
    client, database = persistence_client
    process_id = _create_process(client)

    preview = client.post(
        f"/api/v1/processes/{process_id}/modify/preview",
        json={
            "request_id": "preview-1",
            "base_revision": 0,
            "instruction": "在采购申请后增加供应商确认",
            "locale": "zh-CN",
        },
    )
    assert preview.status_code == 200, preview.text
    payload = preview.json()
    assert payload["base_revision"] == 0
    assert payload["result_revision"] == 1
    assert payload["applied_patch"]["operations"]

    current = client.get(f"/api/v1/processes/{process_id}")
    assert current.status_code == 200
    assert current.json()["current_revision"] == 0
    with database.session_factory() as session:
        revision_count = session.scalar(
            select(func.count()).select_from(ProcessRevisionRecord).where(
                ProcessRevisionRecord.process_id == process_id
            )
        )
        assert revision_count == 1


def test_external_model_policy_requires_admin_opt_in_and_is_audited(persistence_client):
    client, database = persistence_client
    rejected_opt_in = client.post(
        "/api/v1/projects",
        json={"name": "绕过审计项目", "external_model_enabled": True},
    )
    assert rejected_opt_in.status_code == 422

    created = client.post("/api/v1/projects", json={"name": "模型策略项目"})
    assert created.status_code == 201, created.text
    project = created.json()
    project_id = project["id"]
    assert project["external_model_enabled"] is False

    for user_id, role in (
        ("viewer-user", "viewer"),
        ("editor-user", "editor"),
        ("approver-user", "consultant_approver"),
    ):
        response = client.post(
            f"/api/v1/projects/{project_id}/members",
            json={"user_id": user_id, "role": role},
        )
        assert response.status_code == 201, response.text
        headers = {"X-User-ID": user_id}
        forbidden_update = client.put(
            f"/api/v1/projects/{project_id}",
            json={"external_model_enabled": True},
            headers=headers,
        )
        assert forbidden_update.status_code == 403
        forbidden_audit = client.get(
            f"/api/v1/projects/{project_id}/audits",
            headers=headers,
        )
        assert forbidden_audit.status_code == 403

    outsider_headers = {"X-User-ID": "outsider-user"}
    assert (
        client.put(
            f"/api/v1/projects/{project_id}",
            json={"external_model_enabled": True},
            headers=outsider_headers,
        ).status_code
        == 404
    )

    enabled = client.put(
        f"/api/v1/projects/{project_id}",
        json={"external_model_enabled": True},
    )
    assert enabled.status_code == 200, enabled.text
    assert enabled.json()["external_model_enabled"] is True

    unchanged = client.put(
        f"/api/v1/projects/{project_id}",
        json={"external_model_enabled": True},
    )
    assert unchanged.status_code == 200

    disabled = client.put(
        f"/api/v1/projects/{project_id}",
        json={"external_model_enabled": False},
    )
    assert disabled.status_code == 200
    assert disabled.json()["external_model_enabled"] is False

    audits = client.get(f"/api/v1/projects/{project_id}/audits")
    assert audits.status_code == 200, audits.text
    policy_audits = [
        item for item in audits.json() if item["action"] == "external_model_policy_updated"
    ]
    assert [item["after_value"] for item in policy_audits] == [
        {"external_model_enabled": False},
        {"external_model_enabled": True},
    ]
    assert all(item["actor_user_id"] == "local-user" for item in policy_audits)

    session = database.session_factory()
    try:
        audit_records = session.scalars(select(ProjectAuditRecord)).all()
        assert len([item for item in audit_records if item.action == "external_model_policy_updated"]) == 2
    finally:
        session.close()


def test_project_member_changes_are_atomic_and_audited(persistence_client):
    client, database = persistence_client
    created = client.post("/api/v1/projects", json={"name": "成员审计验收项目"})
    assert created.status_code == 201, created.text
    project_id = created.json()["id"]

    added = client.post(
        f"/api/v1/projects/{project_id}/members",
        json={"user_id": "audited-user", "role": "viewer"},
    )
    assert added.status_code == 201, added.text

    unchanged = client.put(
        f"/api/v1/projects/{project_id}/members/audited-user",
        json={"user_id": "audited-user", "role": "viewer"},
    )
    assert unchanged.status_code == 200, unchanged.text

    updated = client.put(
        f"/api/v1/projects/{project_id}/members/audited-user",
        json={"user_id": "audited-user", "role": "editor"},
    )
    assert updated.status_code == 200, updated.text
    assert updated.json()["role"] == "editor"

    removed = client.delete(
        f"/api/v1/projects/{project_id}/members/audited-user"
    )
    assert removed.status_code == 204, removed.text

    audits = client.get(f"/api/v1/projects/{project_id}/audits")
    assert audits.status_code == 200, audits.text
    payload = audits.json()
    assert [item["action"] for item in payload] == [
        "project_member_removed",
        "project_member_role_updated",
        "project_member_added",
    ]
    assert payload[0]["before_value"] == {"user_id": "audited-user", "role": "editor"}
    assert payload[0]["after_value"] == {"user_id": "audited-user", "role": None}
    assert payload[1]["before_value"] == {"user_id": "audited-user", "role": "viewer"}
    assert payload[1]["after_value"] == {"user_id": "audited-user", "role": "editor"}
    assert payload[2]["before_value"] == {"user_id": "audited-user", "role": None}
    assert payload[2]["after_value"] == {"user_id": "audited-user", "role": "viewer"}
    assert all(item["actor_user_id"] == "local-user" for item in payload)

    session = database.session_factory()
    try:
        assert session.get(ProjectMemberRecord, (project_id, "audited-user")) is None
        assert len(session.scalars(select(ProjectAuditRecord)).all()) == 3
    finally:
        session.close()


def test_external_provider_is_blocked_by_default_and_sensitive_logs_are_redacted(
    persistence_client,
):
    client, database = persistence_client
    provider = TrackingExternalProvider()
    app.dependency_overrides[get_provider] = lambda: provider
    try:
        project_response = client.post(
            "/api/v1/projects",
            json={"name": "敏感数据项目", "customer_name": "机密客户"},
        )
        assert project_response.status_code == 201, project_response.text
        project_id = project_response.json()["id"]
        process_response = client.post(
            f"/api/v1/projects/{project_id}/processes",
            json={"name": "脱敏验收流程"},
        )
        assert process_response.status_code == 201, process_response.text
        process_id = process_response.json()["id"]

        blocked = client.post(
            f"/api/v1/processes/{process_id}/modify",
            json={
                "request_id": "external-blocked",
                "base_revision": 0,
                "instruction": "创建流程，联系人：李雷，电话 13800138000",
            },
        )
        assert blocked.status_code == 200, blocked.text
        assert provider.calls == 0
        assert blocked.json()["metrics"]["provider"] == "local"
        assert any("未启用外部模型" in item for item in blocked.json()["warnings"])

        enabled = client.put(
            f"/api/v1/projects/{project_id}",
            json={"external_model_enabled": True},
        )
        assert enabled.status_code == 200, enabled.text
        external = client.post(
            f"/api/v1/processes/{process_id}/modify",
            json={
                "request_id": "external-enabled",
                "base_revision": 1,
                "instruction": (
                    "创建采购订单 ME21N，联系供应商：华东机密供应商，联系人：张三，"
                    "邮箱 secret@example.com，电话 13800138000，金额 100万元"
                ),
            },
        )
        assert external.status_code == 200, external.text
        assert provider.calls == 1
        assert provider.evidence
        assert provider.evidence[0].source_id == "kb-mm-j45-project-seed"
        assert external.json()["metrics"]["provider"] == "fake-external"
        assert not any("未启用外部模型" in item for item in external.json()["warnings"])

        session = database.session_factory()
        try:
            logs = session.scalars(select(ChangeLogRecord)).all()
            log_text = " ".join(
                " ".join(
                    (
                        item.user_prompt,
                        item.decision_summary,
                        str(item.normalized_patch),
                    )
                )
                for item in logs
            )
            for sensitive_value in (
                "华东机密供应商",
                "张三",
                "secret@example.com",
                "13800138000",
                "100万元",
            ):
                assert sensitive_value not in log_text
            assert "[EMAIL_" in log_text
            assert "[PHONE_" in log_text
        finally:
            session.close()
    finally:
        app.dependency_overrides.pop(get_provider, None)


def test_model_cache_reuses_provider_result_after_database_rollback(
    persistence_client,
):
    client, database = persistence_client
    process_id = _create_process(client)
    session = database.session_factory()
    try:
        process = session.get(ProcessRecord, process_id)
        assert process is not None
        project_id = process.project_id
    finally:
        session.close()

    enabled = client.put(
        f"/api/v1/projects/{project_id}",
        json={"external_model_enabled": True},
    )
    assert enabled.status_code == 200, enabled.text

    provider = TrackingExternalProvider()
    cache = ModelResultCache(ttl_seconds=60, max_entries=8)
    app.dependency_overrides[get_provider] = lambda: provider
    app.dependency_overrides[get_model_result_cache] = lambda: cache

    def fail_change_log_insert(*_args) -> None:
        raise OperationalError(
            "INSERT INTO change_log",
            {},
            RuntimeError("simulated cache retry database failure"),
        )

    try:
        event.listen(ChangeLogRecord, "before_insert", fail_change_log_insert)
        try:
            failed = client.post(
                f"/api/v1/processes/{process_id}/modify",
                json={
                    "request_id": "model-cache-db-failure",
                    "base_revision": 0,
                    "instruction": "增加一个外部缓存测试步骤",
                },
            )
        finally:
            event.remove(ChangeLogRecord, "before_insert", fail_change_log_insert)

        assert failed.status_code == 503, failed.text
        assert failed.json()["error"]["code"] == "DATABASE_WRITE_FAILED"
        assert provider.calls == 1

        retried = client.post(
            f"/api/v1/processes/{process_id}/modify",
            json={
                "request_id": "model-cache-db-retry",
                "base_revision": 0,
                "instruction": "增加一个外部缓存测试步骤",
            },
        )
        assert retried.status_code == 200, retried.text
        assert retried.json()["result_revision"] == 1
        assert retried.json()["metrics"]["cache_status"] == "hit"
        assert retried.json()["metrics"]["model_calls"] == 0
        assert retried.json()["metrics"]["attempts"] == 0
        assert provider.calls == 1
    finally:
        app.dependency_overrides.pop(get_provider, None)
        app.dependency_overrides.pop(get_model_result_cache, None)

    session = database.session_factory()
    try:
        process = session.get(ProcessRecord, process_id)
        assert process is not None
        assert process.current_revision == 1
        assert session.scalar(select(func.count()).select_from(ChangeLogRecord)) == 1
    finally:
        session.close()


def test_ungrounded_provider_patch_is_rejected_without_creating_revision(
    persistence_client,
):
    client, database = persistence_client
    process_id = _create_process(client)
    app.dependency_overrides[get_provider] = lambda: UngroundedProvider()
    try:
        response = client.post(
            f"/api/v1/processes/{process_id}/modify",
            json={
                "request_id": "ungrounded-patch",
                "base_revision": 0,
                "instruction": "增加一个虚构交易码步骤",
            },
        )
        assert response.status_code == 422, response.text
        assert response.json()["error"]["code"] == "PATCH_EVIDENCE_INVALID"

        process = client.get(f"/api/v1/processes/{process_id}")
        assert process.status_code == 200
        assert process.json()["current_revision"] == 0
        session = database.session_factory()
        try:
            assert session.scalars(select(ChangeLogRecord)).all() == []
        finally:
            session.close()
    finally:
        app.dependency_overrides.pop(get_provider, None)


def test_local_schema_backfills_legacy_project_owner_membership(persistence_client):
    client, database = persistence_client
    session = database.session_factory()
    try:
        session.add(
            ProjectRecord(
                id="legacy-project",
                tenant_id="local",
                name="旧版开发项目",
                customer_name=None,
                sap_context=SapContext().model_dump(mode="json"),
                created_by="legacy-owner",
            )
        )
        session.commit()
    finally:
        session.close()

    response = client.get("/api/v1/projects", headers={"X-User-ID": "legacy-owner"})
    assert response.status_code == 200
    assert response.json()[0]["id"] == "legacy-project"
    assert response.json()[0]["current_role"] == "project_admin"

    session = database.session_factory()
    try:
        member = session.execute(select(ProjectMemberRecord)).scalar_one()
        assert member.user_id == "legacy-owner"
        assert member.role == "project_admin"
    finally:
        session.close()


def test_persisted_modify_revision_release_and_changelog(persistence_client):
    client, database = persistence_client
    process_id = _create_process(client)

    modify_response = client.post(
        f"/api/v1/processes/{process_id}/modify",
        json={
            "request_id": "persisted-modify-1",
            "base_revision": 0,
            "instruction": "创建直接物料 P2P 流程",
            "evidence_refs": ["kb-mm-j45#standard-steps"],
        },
    )
    assert modify_response.status_code == 200, modify_response.text
    payload = modify_response.json()
    assert payload["base_revision"] == 0
    assert payload["result_revision"] == 1
    assert payload["graph"]["schema_version"] == "2.0"
    assert payload["graph"]["nodes"]
    assert payload["evidence"]
    assert any(item["source_id"] == "kb-mm-j45-project-seed" for item in payload["evidence"])
    purchase_order = next(
        node for node in payload["graph"]["nodes"] if node["label"] == "创建采购订单"
    )
    assert purchase_order["sap"]["tcodes"][0]["status"] == "pending_confirmation"
    assert purchase_order["sap"]["tcodes"][0]["evidence_ref"].startswith(
        "kb-mm-j45-project-seed#"
    )

    markdown_export = client.post(
        f"/api/v1/processes/{process_id}/exports",
        json={"revision_no": 1, "format": "markdown"},
    )
    assert markdown_export.status_code == 200, markdown_export.text
    assert "## 3. 流程步骤" in markdown_export.text
    assert "ME21N" in markdown_export.text
    markdown_disposition = markdown_export.headers["Content-Disposition"]
    assert 'filename="SAP-Blueprint-process_' in markdown_disposition
    markdown_filename = unquote(markdown_disposition.split("filename*=UTF-8''", 1)[1])
    assert re.fullmatch(
        r"SAP-Blueprint-直接物料 P2P-r1-\d{8}-\d{6}\.md",
        markdown_filename,
    )

    docx_export = client.post(
        f"/api/v1/processes/{process_id}/exports",
        json={"revision_no": 1, "format": "docx"},
    )
    assert docx_export.status_code == 200, docx_export.text
    assert docx_export.content.startswith(b"PK")
    document = Document(BytesIO(docx_export.content))
    assert len(document.sections) == 3
    assert [section.orientation for section in document.sections] == [
        WD_ORIENT.PORTRAIT,
        WD_ORIENT.LANDSCAPE,
        WD_ORIENT.PORTRAIT,
    ]
    assert len(document.inline_shapes) == 1
    assert document.inline_shapes[0].width == Inches(9.5)
    document_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "流程步骤" in document_text
    assert any("ME21N" in cell.text for table in document.tables for row in table.rows for cell in row.cells)
    steps_table = next(
        table
        for table in document.tables
        if any(cell.text == "步骤" for cell in table.rows[0].cells)
    )
    assert len(steps_table.rows) - 1 == len(payload["graph"]["nodes"])
    assert all(
        row._tr.get_or_add_trPr().find(qn("w:cantSplit")) is not None
        for table in document.tables
        for row in table.rows
    )
    word_cells = "\n".join(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )
    for node in payload["graph"]["nodes"]:
        assert node["label"] in markdown_export.text
        assert node["label"] in word_cells
    assert "修订号：1" in markdown_export.text
    assert any(
        row.cells[0].text == "修订号" and row.cells[1].text == "1"
        for table in document.tables
        for row in table.rows
    )
    assert "当前修订没有 GAP 候选或正式 GAP。" in markdown_export.text
    assert "当前修订没有 GAP 候选或正式 GAP。" in document_text
    docx_disposition = docx_export.headers["Content-Disposition"]
    docx_filename = unquote(docx_disposition.split("filename*=UTF-8''", 1)[1])
    assert re.fullmatch(
        r"SAP-Blueprint-直接物料 P2P-r1-\d{8}-\d{6}\.docx",
        docx_filename,
    )

    revisions = client.get(f"/api/v1/processes/{process_id}/revisions")
    assert revisions.status_code == 200, revisions.text
    assert [item["revision_no"] for item in revisions.json()] == [1, 0]

    revision_detail = client.get(f"/api/v1/processes/{process_id}/revisions/1")
    assert revision_detail.status_code == 200, revision_detail.text
    assert revision_detail.json()["graph"] == payload["graph"]

    session = database.session_factory()
    try:
        changelog = session.execute(select(ChangeLogRecord)).scalar_one()
        assert changelog.decision_summary
        assert "kb-mm-j45#standard-steps" in changelog.evidence_refs
        assert any(
            item.startswith("kb-mm-j45-project-seed#")
            for item in changelog.evidence_refs
        )
        assert not hasattr(changelog, "ai_reasoning")
    finally:
        session.close()

    release_response = client.post(
        f"/api/v1/processes/{process_id}/releases",
        json={"base_revision": 1},
    )
    assert release_response.status_code == 201, release_response.text
    assert release_response.json()["release_no"] == 1
    assert release_response.json()["lifecycle_state"] == "APPROVED"

    release_detail = client.get(f"/api/v1/processes/{process_id}/releases/1")
    assert release_detail.status_code == 200, release_detail.text
    assert release_detail.json()["release_no"] == 1

    double_release = client.post(
        f"/api/v1/processes/{process_id}/releases",
        json={"base_revision": 1},
    )
    assert double_release.status_code == 409
    assert double_release.json()["error"]["code"] == "REVISION_ALREADY_RELEASED"

    draft = client.post(
        f"/api/v1/processes/{process_id}/drafts",
        json={"source_revision": 1},
    )
    assert draft.status_code == 201, draft.text
    assert draft.json()["revision_no"] == 2
    assert draft.json()["release_no"] is None
    assert draft.json()["lifecycle_state"] == "DRAFT"
    assert draft.json()["graph"]["version"] == 2

    released_snapshot = client.get(f"/api/v1/processes/{process_id}/releases/1")
    assert released_snapshot.status_code == 200
    assert released_snapshot.json()["graph"]["version"] == 1


def test_persisted_diagram_edits_skip_unavailable_knowledge(persistence_client):
    client, database = persistence_client
    process_id = _create_process(client)
    created = client.post(
        f"/api/v1/processes/{process_id}/modify",
        json={
            "request_id": "diagram-routing-seed",
            "base_revision": 0,
            "instruction": "创建直接物料 P2P 流程",
        },
    )
    assert created.status_code == 200, created.text
    node_count = len(created.json()["graph"]["nodes"])

    app.dependency_overrides[get_knowledge_service] = lambda: ExplodingKnowledgeService()
    try:
        instructions = (
            "增加一个合规泳道",
            "把采购申请审批到创建采购订单的连线标签改为已批准",
            "给创建采购申请增加文件图标",
        )
        responses = []
        for revision, instruction in enumerate(instructions, start=1):
            response = client.post(
                f"/api/v1/processes/{process_id}/modify",
                json={
                    "request_id": f"diagram-routing-{revision}",
                    "base_revision": revision,
                    "instruction": instruction,
                },
            )
            assert response.status_code == 200, response.text
            assert response.json()["result_revision"] == revision + 1
            assert len(response.json()["graph"]["nodes"]) == node_count
            responses.append(response.json())
    finally:
        app.dependency_overrides.pop(get_knowledge_service, None)

    assert any(lane["label"] == "合规" for lane in responses[0]["graph"]["lanes"])
    source_ids = {
        node["label"]: node["id"] for node in responses[1]["graph"]["nodes"]
    }
    assert any(
        edge["source"] == source_ids["采购申请审批"]
        and edge["target"] == source_ids["创建采购订单"]
        and edge["label"] == "已批准"
        for edge in responses[1]["graph"]["edges"]
    )
    purchase_request = next(
        node for node in responses[2]["graph"]["nodes"] if node["label"] == "创建采购申请"
    )
    assert purchase_request["icon"] == "file-text"

    session = database.session_factory()
    try:
        assert session.scalar(select(func.count(ChangeLogRecord.id))) == 4
    finally:
        session.close()


def test_persisted_modify_rejects_stale_revision(persistence_client):
    client, _ = persistence_client
    process_id = _create_process(client)

    first = client.post(
        f"/api/v1/processes/{process_id}/modify",
        json={
            "request_id": "persisted-modify-first",
            "base_revision": 0,
            "instruction": "增加一个财务泳道",
        },
    )
    assert first.status_code == 200, first.text

    stale = client.post(
        f"/api/v1/processes/{process_id}/modify",
        json={
            "request_id": "persisted-modify-stale",
            "base_revision": 0,
            "instruction": "增加一个采购泳道",
        },
    )
    assert stale.status_code == 409, stale.text
    assert stale.json()["error"]["code"] == "REVISION_CONFLICT"
    assert stale.json()["error"]["details"] == {
        "expected_revision": 0,
        "current_revision": 1,
    }


def test_database_write_failure_rolls_back_process_revision_and_changelog(
    persistence_client,
):
    client, database = persistence_client
    process_id = _create_process(client)
    process_response = client.get(f"/api/v1/processes/{process_id}")
    graph = process_response.json()["graph"]
    graph["version"] = 1
    graph["title"] = "不应保存的流程标题"

    def fail_change_log_insert(*_args) -> None:
        raise OperationalError(
            "INSERT INTO change_log",
            {},
            RuntimeError("simulated database detail must stay private"),
        )

    event.listen(ChangeLogRecord, "before_insert", fail_change_log_insert)
    try:
        response = client.post(
            f"/api/v1/processes/{process_id}/save",
            headers={"X-Request-ID": "database-failure-rollback"},
            json={
                "request_id": "database-failure-rollback",
                "base_revision": 0,
                "graph": graph,
                "summary": "触发事务回滚",
            },
        )
    finally:
        event.remove(ChangeLogRecord, "before_insert", fail_change_log_insert)

    assert response.status_code == 503, response.text
    assert response.headers["X-Request-ID"] == "database-failure-rollback"
    assert response.json()["error"] == {
        "code": "DATABASE_WRITE_FAILED",
        "message": "数据库写入失败，未保存任何更改，请稍后重试。",
        "request_id": "database-failure-rollback",
        "details": {},
    }
    assert "simulated database detail" not in response.text

    session = database.session_factory()
    try:
        process = session.get(ProcessRecord, process_id)
        assert process is not None
        assert process.current_revision == 0
        assert process.status == "DRAFT"
        assert session.scalar(
            select(func.count()).select_from(ProcessRevisionRecord)
        ) == 1
        assert session.scalar(select(func.count()).select_from(ChangeLogRecord)) == 0
    finally:
        session.close()


def test_dependency_or_export_failures_keep_current_graph_available(
    persistence_client, monkeypatch
):
    client, database = persistence_client
    process_id = _create_process(client)

    app.dependency_overrides[get_provider] = lambda: ExplodingProvider()
    try:
        provider_failure = client.post(
            f"/api/v1/processes/{process_id}/modify",
            json={
                "request_id": "provider-failure-persisted",
                "base_revision": 0,
                "instruction": "增加一个审批步骤",
            },
        )
    finally:
        app.dependency_overrides.pop(get_provider, None)
    assert provider_failure.status_code == 502, provider_failure.text
    assert provider_failure.json()["error"]["code"] == "PROVIDER_UNAVAILABLE"

    app.dependency_overrides[get_knowledge_service] = lambda: ExplodingKnowledgeService()
    try:
        knowledge_failure = client.post(
            f"/api/v1/processes/{process_id}/modify",
            json={
                "request_id": "knowledge-failure-persisted",
                "base_revision": 0,
                "instruction": "增加一个审批步骤",
            },
        )
    finally:
        app.dependency_overrides.pop(get_knowledge_service, None)
    assert knowledge_failure.status_code == 503, knowledge_failure.text
    assert knowledge_failure.json()["error"]["code"] == "KNOWLEDGE_UNAVAILABLE"

    def fail_docx_render(_model):
        raise RuntimeError("private export rendering detail")

    monkeypatch.setattr(export_service, "render_docx", fail_docx_render)
    export_failure = client.post(
        f"/api/v1/processes/{process_id}/exports",
        json={"revision_no": 0, "format": "docx"},
    )
    assert export_failure.status_code == 500, export_failure.text
    assert export_failure.json()["error"]["code"] == "INTERNAL_SERVER_ERROR"
    assert "private export rendering detail" not in export_failure.text

    current = client.get(f"/api/v1/processes/{process_id}")
    assert current.status_code == 200, current.text
    assert current.json()["current_revision"] == 0
    assert current.json()["graph"]["version"] == 0

    session = database.session_factory()
    try:
        assert session.scalar(select(func.count()).select_from(ProcessRevisionRecord)) == 1
        assert session.scalar(select(func.count()).select_from(ChangeLogRecord)) == 0
    finally:
        session.close()


def test_gap_decision_is_audited_and_requires_valid_lifecycle(persistence_client):
    client, database = persistence_client
    graph = {
        "schema_version": "2.0",
        "graph_id": "gap-graph",
        "version": 0,
        "title": "GAP 流程",
        "module": "MM",
        "process_scope": "P2P",
        "sap_context": {},
        "direction": "TB",
        "nodes": [
            {
                "id": "gap-node",
                "type": "task",
                "label": "审批",
                "sap": {
                    "gap": {
                        "status": "candidate",
                        "description": "动态审批",
                        "evidence_refs": ["kb-gap"],
                    }
                },
            }
        ],
        "edges": [],
        "lanes": [],
        "layout": {},
    }
    gap_project = client.post(
        "/api/v1/projects", json={"name": "GAP 决策项目"}
    )
    assert gap_project.status_code == 201, gap_project.text
    project_id = gap_project.json()["id"]
    process_response = client.post(
        f"/api/v1/projects/{project_id}/processes",
        json={"name": "GAP 流程 2", "initial_graph": graph},
    )
    assert process_response.status_code == 201, process_response.text
    process_id = process_response.json()["id"]
    decision = client.post(
        f"/api/v1/processes/{process_id}/gaps/gap-node/decisions",
        json={"base_revision": 0, "to_status": "confirmed", "comment": "顾问确认差异"},
    )
    assert decision.status_code == 200, decision.text
    assert decision.json()["to_status"] == "confirmed"
    assert decision.json()["graph"]["nodes"][0]["sap"]["gap"]["status"] == "confirmed"

    session = database.session_factory()
    try:
        audited = session.execute(select(GapDecisionRecord)).scalar_one()
        assert audited.comment == "顾问确认差异"
        assert audited.decided_by == "local-user"
        human_change = session.execute(
            select(ChangeLogRecord).where(ChangeLogRecord.provider == "human")
        ).scalar_one()
        assert human_change.model == "manual"
    finally:
        session.close()

    invalid = client.post(
        f"/api/v1/processes/{process_id}/gaps/gap-node/decisions",
        json={"base_revision": 1, "to_status": "confirmed", "comment": "重复确认"},
    )
    assert invalid.status_code == 409
    assert invalid.json()["error"]["code"] == "GAP_INVALID_TRANSITION"


def test_project_roles_are_loaded_from_membership_and_isolate_projects(persistence_client):
    client, _ = persistence_client
    project_response = client.post("/api/v1/projects", json={"name": "权限验收项目"})
    assert project_response.status_code == 201, project_response.text
    project_id = project_response.json()["id"]
    assert project_response.json()["current_role"] == "project_admin"

    process_response = client.post(
        f"/api/v1/projects/{project_id}/processes",
        json={"name": "权限验收流程"},
    )
    assert process_response.status_code == 201, process_response.text
    process_id = process_response.json()["id"]

    add_viewer = client.post(
        f"/api/v1/projects/{project_id}/members",
        json={"user_id": "viewer-user", "role": "viewer"},
    )
    assert add_viewer.status_code == 201, add_viewer.text
    assert add_viewer.json()["role"] == "viewer"

    spoofed_headers = {
        "X-User-ID": "viewer-user",
        "X-Project-Role": "project_admin",
    }
    visible_projects = client.get("/api/v1/projects", headers=spoofed_headers)
    assert visible_projects.status_code == 200
    assert visible_projects.json()[0]["current_role"] == "viewer"
    assert client.get(f"/api/v1/processes/{process_id}", headers=spoofed_headers).status_code == 200

    forbidden_create = client.post(
        f"/api/v1/projects/{project_id}/processes",
        json={"name": "越权流程"},
        headers=spoofed_headers,
    )
    assert forbidden_create.status_code == 403
    assert forbidden_create.json()["error"]["code"] == "PROJECT_PERMISSION_DENIED"
    forbidden_members = client.get(
        f"/api/v1/projects/{project_id}/members", headers=spoofed_headers
    )
    assert forbidden_members.status_code == 403
    knowledge_request = {
        "module": "MM",
        "process_scope": "P2P",
        "query": "采购申请审批",
        "sap_context": {},
        "top_k": 3,
    }
    forbidden_knowledge = client.post(
        f"/api/v1/projects/{project_id}/knowledge/search",
        json=knowledge_request,
        headers=spoofed_headers,
    )
    assert forbidden_knowledge.status_code == 403

    outsider_headers = {"X-User-ID": "outsider-user"}
    assert client.get("/api/v1/projects", headers=outsider_headers).json() == []
    hidden_process = client.get(
        f"/api/v1/processes/{process_id}", headers=outsider_headers
    )
    assert hidden_process.status_code == 404
    assert hidden_process.json()["error"]["code"] == "PROJECT_NOT_FOUND"
    hidden_export = client.post(
        f"/api/v1/processes/{process_id}/exports",
        json={"revision_no": 0, "format": "markdown"},
        headers=outsider_headers,
    )
    assert hidden_export.status_code == 404
    hidden_knowledge = client.post(
        f"/api/v1/projects/{project_id}/knowledge/search",
        json=knowledge_request,
        headers=outsider_headers,
    )
    assert hidden_knowledge.status_code == 404

    update_viewer = client.put(
        f"/api/v1/projects/{project_id}/members/viewer-user",
        json={"user_id": "viewer-user", "role": "editor"},
    )
    assert update_viewer.status_code == 200
    editor_headers = {"X-User-ID": "viewer-user"}
    editor_create = client.post(
        f"/api/v1/projects/{project_id}/processes",
        json={"name": "编辑者流程"},
        headers=editor_headers,
    )
    assert editor_create.status_code == 201, editor_create.text
    editor_knowledge = client.post(
        f"/api/v1/projects/{project_id}/knowledge/search",
        json=knowledge_request,
        headers=editor_headers,
    )
    assert editor_knowledge.status_code == 200, editor_knowledge.text
    editor_release = client.post(
        f"/api/v1/processes/{process_id}/releases",
        json={"base_revision": 0},
        headers=editor_headers,
    )
    assert editor_release.status_code == 403

    last_admin = client.put(
        f"/api/v1/projects/{project_id}/members/local-user",
        json={"user_id": "local-user", "role": "editor"},
    )
    assert last_admin.status_code == 409
    assert last_admin.json()["error"]["code"] == "LAST_PROJECT_ADMIN"
    remove_last_admin = client.delete(
        f"/api/v1/projects/{project_id}/members/local-user"
    )
    assert remove_last_admin.status_code == 409
    assert remove_last_admin.json()["error"]["code"] == "LAST_PROJECT_ADMIN"


def test_gap_status_requires_decision_and_release_requires_decision_audit(persistence_client):
    client, _ = persistence_client
    project_response = client.post("/api/v1/projects", json={"name": "发布预检项目"})
    project_id = project_response.json()["id"]
    graph = {
        "schema_version": "2.0",
        "graph_id": "release-gap-graph",
        "version": 0,
        "title": "发布 GAP 流程",
        "module": "MM",
        "process_scope": "P2P",
        "sap_context": {},
        "direction": "LR",
        "nodes": [
            {"id": "start", "type": "start", "label": "开始"},
            {
                "id": "approval",
                "type": "task",
                "label": "采购审批",
                "sap": {
                    "gap": {
                        "status": "candidate",
                        "description": "需要定制审批",
                        "evidence_refs": ["kb-gap-approval"],
                    }
                },
            },
            {"id": "end", "type": "end", "label": "结束"},
        ],
        "edges": [
            {"id": "e1", "source": "start", "target": "approval"},
            {"id": "e2", "source": "approval", "target": "end"},
        ],
        "lanes": [],
        "layout": {},
    }
    process_response = client.post(
        f"/api/v1/projects/{project_id}/processes",
        json={"name": "发布 GAP 流程", "initial_graph": graph},
    )
    assert process_response.status_code == 201, process_response.text
    process_id = process_response.json()["id"]

    direct_graph = process_response.json()["graph"]
    direct_graph["version"] = 1
    direct_graph["nodes"][1]["sap"]["gap"]["status"] = "confirmed"
    direct_save = client.post(
        f"/api/v1/processes/{process_id}/save",
        json={
            "request_id": "direct-confirm",
            "base_revision": 0,
            "graph": direct_graph,
            "summary": "直接确认 GAP",
        },
    )
    assert direct_save.status_code == 409
    assert direct_save.json()["error"]["code"] == "GAP_DECISION_REQUIRED"

    decision = client.post(
        f"/api/v1/processes/{process_id}/gaps/approval/decisions",
        json={
            "base_revision": 0,
            "to_status": "confirmed",
            "comment": "顾问确认需要定制审批",
        },
    )
    assert decision.status_code == 200, decision.text
    release = client.post(
        f"/api/v1/processes/{process_id}/releases",
        json={"base_revision": 1},
    )
    assert release.status_code == 201, release.text

    unaudited_graph = graph.copy()
    unaudited_graph["graph_id"] = "unaudited-gap-graph"
    unaudited_graph["nodes"] = [dict(node) for node in graph["nodes"]]
    unaudited_graph["nodes"][1] = {
        **unaudited_graph["nodes"][1],
        "sap": {
            "gap": {
                "status": "confirmed",
                "description": "未经决策直接确认",
                "evidence_refs": ["kb-gap-approval"],
            }
        },
    }
    unaudited_process = client.post(
        f"/api/v1/projects/{project_id}/processes",
        json={"name": "未审计 GAP 流程", "initial_graph": unaudited_graph},
    )
    assert unaudited_process.status_code == 201, unaudited_process.text
    unaudited_release = client.post(
        f"/api/v1/processes/{unaudited_process.json()['id']}/releases",
        json={"base_revision": 0},
    )
    assert unaudited_release.status_code == 422
    assert unaudited_release.json()["error"]["code"] == "RELEASE_PREFLIGHT_FAILED"
    assert unaudited_release.json()["error"]["details"]["gap_without_decision_audit"]


def test_versioned_acceptance_demo_completes_release_and_exports(persistence_client):
    client, _ = persistence_client
    scenario_path = (
        Path(__file__).resolve().parents[3]
        / "examples"
        / "mm-p2p-acceptance-demo.json"
    )
    scenario = json.loads(scenario_path.read_text(encoding="utf-8"))

    project_response = client.post("/api/v1/projects", json=scenario["project"])
    assert project_response.status_code == 201, project_response.text
    project_id = project_response.json()["id"]
    process_response = client.post(
        f"/api/v1/projects/{project_id}/processes",
        json=scenario["process"],
    )
    assert process_response.status_code == 201, process_response.text
    process_id = process_response.json()["id"]
    revision = process_response.json()["current_revision"]
    graph = process_response.json()["graph"]

    for index, instruction in enumerate(scenario["instructions"], start=1):
        modified = client.post(
            f"/api/v1/processes/{process_id}/modify",
            json={
                "request_id": f"acceptance-demo-{index}",
                "base_revision": revision,
                "instruction": instruction,
                "locale": "zh-CN",
            },
        )
        assert modified.status_code == 200, modified.text
        revision = modified.json()["result_revision"]
        graph = modified.json()["graph"]

    expected = scenario["expected"]
    assert len(graph["nodes"]) == expected["node_count"]
    assert len(graph["edges"]) == expected["edge_count"]
    assert len(graph["lanes"]) == expected["lane_count"]
    assert {node["label"] for node in graph["nodes"]} == set(expected["node_labels"])
    assert {lane["label"] for lane in graph["lanes"]} == set(expected["lane_labels"])
    approval_id = next(
        node["id"]
        for node in graph["nodes"]
        if node["label"] == expected["approval_source_label"]
    )
    purchase_order_id = next(
        node["id"]
        for node in graph["nodes"]
        if node["label"] == expected["approval_target_label"]
    )
    approval_edge = next(
        edge
        for edge in graph["edges"]
        if edge["source"] == approval_id and edge["target"] == purchase_order_id
    )
    assert approval_edge["label"] == expected["approval_edge_label"]

    release = client.post(
        f"/api/v1/processes/{process_id}/releases",
        json={"base_revision": revision},
    )
    assert release.status_code == 201, release.text
    assert release.json()["lifecycle_state"] == "APPROVED"

    markdown = client.post(
        f"/api/v1/processes/{process_id}/exports",
        json={"revision_no": revision, "format": "markdown"},
    )
    docx = client.post(
        f"/api/v1/processes/{process_id}/exports",
        json={"revision_no": revision, "format": "docx"},
    )
    assert markdown.status_code == 200, markdown.text
    assert scenario["process"]["name"] in markdown.text
    assert docx.status_code == 200, docx.text
    assert docx.content[:2] == b"PK"


def test_async_export_job_persists_status_download_and_expiration(persistence_client):
    client, database = persistence_client
    process_id = _create_process(client)
    modified = client.post(
        f"/api/v1/processes/{process_id}/modify",
        json={
            "request_id": "async-export-source",
            "base_revision": 0,
            "instruction": "创建直接物料 P2P 流程",
        },
    )
    assert modified.status_code == 200, modified.text

    created = client.post(
        f"/api/v1/processes/{process_id}/exports/jobs",
        json={"revision_no": 1, "format": "markdown"},
    )
    assert created.status_code == 202, created.text
    assert created.json()["status"] == "pending"
    assert created.json()["attempt_count"] == 0
    export_id = created.json()["export_id"]

    status = client.get(f"/api/v1/processes/{process_id}/exports/jobs/{export_id}")
    assert status.status_code == 200, status.text
    payload = status.json()
    assert payload["status"] == "completed"
    assert payload["attempt_count"] == 1
    assert payload["filename"].endswith(".md")
    assert payload["content_length"] > 100
    assert payload["download_url"].endswith(f"/{export_id}/download")

    download = client.get(payload["download_url"])
    assert download.status_code == 200, download.text
    assert "## 3. 流程步骤" in download.text

    session = database.session_factory()
    try:
        job = session.get(ExportJobRecord, export_id)
        assert job is not None
        assert job.content == download.content
        assert job.artifact_backend == "database"
        assert job.content_sha256 is not None
        job.expires_at = datetime.now(UTC) - timedelta(seconds=1)
        session.commit()
    finally:
        session.close()

    expired = client.get(f"/api/v1/processes/{process_id}/exports/jobs/{export_id}")
    assert expired.status_code == 200, expired.text
    assert expired.json()["status"] == "expired"
    assert expired.json()["content_length"] is None
    assert expired.json()["download_url"] is None
    expired_download = client.get(
        f"/api/v1/processes/{process_id}/exports/jobs/{export_id}/download"
    )
    assert expired_download.status_code == 410
    assert expired_download.json()["error"]["code"] == "EXPORT_EXPIRED"


def test_filesystem_export_keeps_binary_out_of_database_and_cleans_expired_artifact(
    persistence_client,
    tmp_path,
):
    client, database = persistence_client
    storage_root = tmp_path / "export-artifacts"
    settings = Settings(
        _env_file=None,
        export_execution_mode="worker",
        export_storage_backend="filesystem",
        export_storage_path=str(storage_root),
    )
    app.dependency_overrides[get_settings] = lambda: settings
    try:
        process_id = _create_process(client)
        created = client.post(
            f"/api/v1/processes/{process_id}/exports/jobs",
            json={"revision_no": 0, "format": "markdown"},
        )
        export_id = created.json()["export_id"]
        assert created.json()["status"] == "pending"
        assert ExportJobWorker(database, settings).run_once() == 1
        status = client.get(f"/api/v1/processes/{process_id}/exports/jobs/{export_id}")
        assert status.status_code == 200, status.text
        assert status.json()["status"] == "completed"

        with database.session_factory() as session:
            job = session.get(ExportJobRecord, export_id)
            assert job is not None
            assert job.content is None
            assert job.artifact_backend == "filesystem"
            assert job.artifact_key is not None
            assert job.content_sha256 is not None
            artifact_path = storage_root / job.artifact_key
            assert artifact_path.is_file()
            expected_content = artifact_path.read_bytes()

        download = client.get(status.json()["download_url"])
        assert download.status_code == 200, download.text
        assert download.content == expected_content

        artifact_path.write_bytes(bytes([expected_content[0] ^ 1]) + expected_content[1:])
        corrupted = client.get(status.json()["download_url"])
        assert corrupted.status_code == 503
        assert corrupted.json()["error"]["code"] == "EXPORT_STORAGE_UNAVAILABLE"
        assert str(storage_root) not in corrupted.text

        with database.session_factory() as session:
            job = session.get(ExportJobRecord, export_id)
            assert job is not None
            job.expires_at = datetime.now(UTC) - timedelta(seconds=1)
            session.commit()

        expired = client.get(f"/api/v1/processes/{process_id}/exports/jobs/{export_id}")
        assert expired.status_code == 200, expired.text
        assert expired.json()["status"] == "expired"
        assert not artifact_path.exists()
        with database.session_factory() as session:
            job = session.get(ExportJobRecord, export_id)
            assert job is not None
            assert job.artifact_key is None
            assert job.content_sha256 is None
    finally:
        app.dependency_overrides.pop(get_settings, None)


def test_async_export_failure_is_safe_and_query_is_process_scoped(
    persistence_client, monkeypatch
):
    client, _ = persistence_client
    process_id = _create_process(client)

    def explode_docx(_model):
        raise RuntimeError("sensitive renderer details")

    monkeypatch.setattr(export_service, "render_docx", explode_docx)
    created = client.post(
        f"/api/v1/processes/{process_id}/exports/jobs",
        json={"revision_no": 0, "format": "docx"},
    )
    assert created.status_code == 202, created.text
    export_id = created.json()["export_id"]

    status = client.get(f"/api/v1/processes/{process_id}/exports/jobs/{export_id}")
    assert status.status_code == 200, status.text
    assert status.json()["status"] == "failed"
    assert status.json()["error_code"] == "EXPORT_RENDER_FAILED"
    assert "sensitive" not in status.text

    failed_download = client.get(
        f"/api/v1/processes/{process_id}/exports/jobs/{export_id}/download"
    )
    assert failed_download.status_code == 409
    assert failed_download.json()["error"]["code"] == "EXPORT_RENDER_FAILED"
    assert "sensitive" not in failed_download.text

    other_process_id = _create_process(client)
    hidden = client.get(
        f"/api/v1/processes/{other_process_id}/exports/jobs/{export_id}"
    )
    assert hidden.status_code == 404
    assert hidden.json()["error"]["code"] == "EXPORT_NOT_FOUND"


def test_stale_async_export_is_reclaimed_on_status_query(persistence_client):
    client, database = persistence_client
    process_id = _create_process(client)
    session = database.session_factory()
    try:
        session.add(
            ExportJobRecord(
                id="export-stale",
                process_id=process_id,
                revision_no=0,
                format="markdown",
                status="running",
                created_by="local-user",
                started_at=datetime.now(UTC) - timedelta(minutes=10),
            )
        )
        session.commit()
    finally:
        session.close()

    first = client.get(
        f"/api/v1/processes/{process_id}/exports/jobs/export-stale"
    )
    assert first.status_code == 200, first.text
    assert first.json()["status"] == "running"

    completed = client.get(
        f"/api/v1/processes/{process_id}/exports/jobs/export-stale"
    )
    assert completed.status_code == 200, completed.text
    assert completed.json()["status"] == "completed"
    assert completed.json()["attempt_count"] == 1
    assert completed.json()["download_url"].endswith("/export-stale/download")


def test_worker_mode_keeps_api_pending_until_worker_claims_job(persistence_client):
    client, database = persistence_client
    process_id = _create_process(client)
    settings = Settings(
        _env_file=None,
        export_execution_mode="worker",
        export_worker_batch_size=4,
        database_url=database.url,
        database_auto_create=False,
    )
    app.dependency_overrides[get_settings] = lambda: settings
    try:
        created = client.post(
            f"/api/v1/processes/{process_id}/exports/jobs",
            json={"revision_no": 0, "format": "markdown"},
        )
        assert created.status_code == 202, created.text
        export_id = created.json()["export_id"]
        assert created.json()["status"] == "pending"
        assert created.json()["attempt_count"] == 0

        polled = client.get(f"/api/v1/processes/{process_id}/exports/jobs/{export_id}")
        assert polled.status_code == 200, polled.text
        assert polled.json()["status"] == "pending"
        assert polled.json()["attempt_count"] == 0

        worker = ExportJobWorker(database, settings)
        assert worker.run_once() == 1
        assert worker.run_once() == 0

        completed = client.get(f"/api/v1/processes/{process_id}/exports/jobs/{export_id}")
        assert completed.status_code == 200, completed.text
        assert completed.json()["status"] == "completed"
        assert completed.json()["attempt_count"] == 1
        assert completed.json()["download_url"].endswith(f"/{export_id}/download")
    finally:
        app.dependency_overrides.pop(get_settings, None)


def test_export_job_lease_fencing_rejects_stale_worker_writes(persistence_client):
    client, database = persistence_client
    process_id = _create_process(client)
    with database.session_factory() as session:
        repository = BlueprintRepository(session)
        job = repository.create_export_job(
            process_id=process_id,
            revision_no=0,
            format="markdown",
            user_id="local-user",
        )
        first_claim = repository.claim_export_job(job.id, stale_minutes=5)
        assert first_claim is not None

        claimed_job = repository.require_export_job_unscoped(job.id)
        claimed_job.started_at = datetime.now(UTC) - timedelta(minutes=10)
        claimed_job.heartbeat_at = datetime.now(UTC) - timedelta(minutes=10)
        session.commit()
        second_claim = repository.claim_export_job(job.id, stale_minutes=5)
        assert second_claim is not None
        assert second_claim != first_claim
        assert not repository.renew_export_job_lease(
            export_id=job.id,
            claim_token=first_claim,
        )
        assert repository.renew_export_job_lease(
            export_id=job.id,
            claim_token=second_claim,
        )

        assert not repository.complete_export_job(
            export_id=job.id,
            claim_token=first_claim,
            filename="stale.md",
            fallback_filename="stale.md",
            media_type="text/markdown",
            content=b"stale",
            retention_hours=24,
        )
        assert not repository.fail_export_job(export_id=job.id, claim_token=first_claim)
        assert repository.complete_export_job(
            export_id=job.id,
            claim_token=second_claim,
            filename="winner.md",
            fallback_filename="winner.md",
            media_type="text/markdown",
            content=b"winner",
            retention_hours=24,
        )
        session.expire_all()
        completed = repository.require_export_job_unscoped(job.id)
        assert completed.status == "completed"
        assert completed.attempt_count == 2
        assert completed.claim_token is None
        assert completed.content == b"winner"


def test_long_export_heartbeat_prevents_false_stale_reclaim(
    persistence_client,
    monkeypatch,
):
    client, database = persistence_client
    process_id = _create_process(client)
    with database.session_factory() as session:
        repository = BlueprintRepository(session)
        job = repository.create_export_job(
            process_id=process_id,
            revision_no=0,
            format="markdown",
            user_id="local-user",
        )
        export_id = job.id

    render_started = Event()
    release_render = Event()
    original_render_export = export_service.render_export

    def slow_render_export(model, format):
        render_started.set()
        if not release_render.wait(timeout=2):
            raise RuntimeError("test did not release long export render")
        return original_render_export(model, format)

    monkeypatch.setattr(export_service, "render_export", slow_render_export)
    outcome: dict[str, object] = {}
    stale_minutes = 0.01
    heartbeat_seconds = 0.05

    def run_export() -> None:
        try:
            outcome["claimed"] = export_service.process_export_job(
                export_id,
                database.engine,
                retention_hours=24,
                stale_minutes=stale_minutes,
                heartbeat_seconds=heartbeat_seconds,
            )
        except BaseException as exc:
            outcome["error"] = exc

    thread = Thread(target=run_export, daemon=True)
    thread.start()
    assert render_started.wait(timeout=2)
    try:
        # Keep the render running beyond the stale window while the heartbeat renews the lease.
        sleep(stale_minutes * 60 + 0.15)
        with database.session_factory() as session:
            repository = BlueprintRepository(session)
            running = repository.require_export_job_unscoped(export_id)
            assert running.status == "running"
            assert running.heartbeat_at is not None
            assert running.started_at is not None
            assert running.heartbeat_at > running.started_at
            assert export_id not in repository.list_export_job_candidates(
                stale_minutes=stale_minutes,
                limit=10,
            )
            assert repository.claim_export_job(export_id, stale_minutes=stale_minutes) is None
    finally:
        release_render.set()
        thread.join(timeout=3)

    assert not thread.is_alive()
    assert "error" not in outcome
    assert outcome["claimed"] is True
    with database.session_factory() as session:
        completed = BlueprintRepository(session).require_export_job_unscoped(export_id)
        assert completed.status == "completed"
        assert completed.attempt_count == 1


def test_lost_export_lease_deletes_attempt_scoped_filesystem_artifact(
    persistence_client,
    tmp_path,
    monkeypatch,
):
    client, database = persistence_client
    process_id = _create_process(client)
    with database.session_factory() as session:
        job = BlueprintRepository(session).create_export_job(
            process_id=process_id,
            revision_no=0,
            format="markdown",
            user_id="local-user",
        )
        export_id = job.id

    original_render_export = export_service.render_export
    replacement_claim = "export-claim-replacement"

    def render_then_replace_lease(model, format):
        rendered = original_render_export(model, format)
        with database.session_factory() as session:
            claimed_job = session.get(ExportJobRecord, export_id)
            assert claimed_job is not None
            claimed_job.claim_token = replacement_claim
            session.commit()
        return rendered

    monkeypatch.setattr(export_service, "render_export", render_then_replace_lease)
    storage_root = tmp_path / "lease-artifacts"
    store = FilesystemArtifactStore(str(storage_root))

    assert export_service.process_export_job(
        export_id,
        database.engine,
        retention_hours=24,
        stale_minutes=5,
        heartbeat_seconds=30,
        artifact_store=store,
    )

    assert list(storage_root.rglob("*.artifact")) == []
    with database.session_factory() as session:
        running = session.get(ExportJobRecord, export_id)
        assert running is not None
        assert running.status == "running"
        assert running.claim_token == replacement_claim
        assert running.artifact_key is None
