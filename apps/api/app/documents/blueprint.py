from collections.abc import Iterable
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

from app.models.graph import GraphDocument


@dataclass(frozen=True)
class BlueprintDocumentModel:
    project_name: str
    customer_name: str | None
    process_name: str
    process_id: str
    revision_no: int
    release_no: int | None
    lifecycle_state: str
    created_by: str
    created_at: str
    graph: GraphDocument


def render_markdown(model: BlueprintDocumentModel) -> str:
    graph = model.graph
    version_label = (
        f"发布版本 {model.release_no} / 修订 {model.revision_no}"
        if model.release_no is not None
        else f"草稿修订 {model.revision_no}"
    )
    lines = [
        f"# {graph.title}",
        "",
        f"> {model.project_name} | {version_label} | {model.lifecycle_state}",
        "",
        "## 1. 项目与 SAP 环境",
        "",
        "| 字段 | 内容 |",
        "| --- | --- |",
        f"| 项目 | {_md(model.project_name)} |",
        f"| 客户 | {_md(model.customer_name or '未设置')} |",
        f"| 流程 | {_md(model.process_name)} |",
        f"| 模块 / 范围 | {_md(graph.module)} / {_md(graph.process_scope)} |",
        f"| SAP | {_md(graph.sap_context.edition)} { _md(graph.sap_context.release)} |",
        f"| 部署 / 国家 | {_md(graph.sap_context.deployment)} / {_md(graph.sap_context.country)} |",
        f"| 版本 | {_md(version_label)} |",
        "",
        "## 2. 流程图",
        "",
        *_mermaid_lines(graph),
        "",
        "## 3. 流程步骤",
        "",
        "| # | 泳道 | 步骤 | 类型 | T-Code | 角色 | 元数据状态 |",
        "| ---: | --- | --- | --- | --- | --- | --- |",
    ]
    lane_labels = {lane.id: lane.label for lane in graph.lanes}
    for index, node in enumerate(graph.nodes, 1):
        lines.append(
            "| "
            + " | ".join(
                [
                    str(index),
                    _md(lane_labels.get(node.lane_id or "", "未分配")),
                    _md(node.label),
                    _md(node.sap.step_type or node.type),
                    _md(", ".join(item.code for item in node.sap.tcodes) or "-"),
                    _md(", ".join(node.sap.roles) or "-"),
                    _md(_metadata_summary(node)),
                ]
            )
            + " |"
        )

    lines.extend(["", "## 4. SAP 元数据与证据", ""])
    evidence_rows = list(_evidence_rows(graph))
    if evidence_rows:
        lines.extend(
            [
                "| 步骤 | 对象 | 状态 | 证据引用 |",
                "| --- | --- | --- | --- |",
                *[
                    f"| {_md(step)} | {_md(item)} | {_md(status)} | `{_md(evidence)}` |"
                    for step, item, status, evidence in evidence_rows
                ],
            ]
        )
    else:
        lines.append("当前修订尚无 SAP 元数据证据。")

    lines.extend(["", "## 5. GAP List", ""])
    gaps = [(node.label, node.sap.gap) for node in graph.nodes if node.sap.gap.status != "none"]
    if gaps:
        lines.extend(
            [
                "| 步骤 | 状态 | 分类 | 差异描述 | 建议方案 | 负责人 | 证据 |",
                "| --- | --- | --- | --- | --- | --- | --- |",
            ]
        )
        for label, gap in gaps:
            lines.append(
                "| "
                + " | ".join(
                    [
                        _md(label),
                        _md(_gap_status(gap.status)),
                        _md(gap.category or "待分类"),
                        _md(gap.description or "待补充"),
                        _md(gap.recommendation or "待评估"),
                        _md(gap.owner or "未分配"),
                        _md(", ".join(gap.evidence_refs) or "待确认"),
                    ]
                )
                + " |"
            )
    else:
        lines.append("当前修订没有 GAP 候选或正式 GAP。")

    lines.extend(
        [
            "",
            "## 6. 版本信息",
            "",
            f"- 流程 ID：`{_md(model.process_id)}`",
            f"- 修订号：{model.revision_no}",
            f"- 发布版本：{model.release_no if model.release_no is not None else '未发布'}",
            f"- 生命周期：{_md(model.lifecycle_state)}",
            f"- 创建人：{_md(model.created_by)}",
            f"- 创建时间：{_md(model.created_at)}",
            "",
            "> 待确认字段和候选 GAP 必须由 SAP 顾问复核；本文档不会把模型建议自动升级为已验证结论。",
            "",
        ]
    )
    return "\n".join(lines)


def render_docx(model: BlueprintDocumentModel) -> bytes:
    document = Document()
    _configure_document(document, model)
    _add_title_block(document, model)
    _add_project_section(document, model)
    _add_heading(document, "2. 流程图", level=1)
    diagram = render_graph_png(model.graph)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(8)
    picture = paragraph.add_run().add_picture(diagram, width=Inches(6.5))
    picture._inline.docPr.set("descr", "SAP 业务流程及泳道图")
    picture._inline.docPr.set("title", "业务流程图")
    caption = document.add_paragraph("图 1  业务流程及泳道")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(10)
    for run in caption.runs:
        _set_run_font(run, size=9, color="68716B")

    _add_steps_section(document, model.graph)
    _add_evidence_section(document, model.graph)
    _add_gaps_section(document, model.graph)
    _add_version_section(document, model)

    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def render_graph_png(graph: GraphDocument) -> BytesIO:
    lanes = list(graph.lanes)
    unassigned = [node for node in graph.nodes if not node.lane_id]
    lane_entries: list[tuple[str, str, str]] = [
        (lane.id, lane.label, lane.color) for lane in lanes
    ]
    if unassigned or not lane_entries:
        lane_entries.append(("__unassigned__", "未分配", "#7B817D"))

    grouped = {
        lane_id: sorted(
            [node for node in graph.nodes if (node.lane_id or "__unassigned__") == lane_id],
            key=lambda node: (
                graph.layout.get(node.id).x if node.id in graph.layout else graph.nodes.index(node)
            ),
        )
        for lane_id, _, _ in lane_entries
    }
    node_order = {node.id: index for index, node in enumerate(graph.nodes)}
    width = max(1400, 300 + max(len(graph.nodes), 1) * 230)
    height = max(420, 86 + len(lane_entries) * 178)
    image = Image.new("RGB", (width, height), "#F7F7F4")
    draw = ImageDraw.Draw(image)
    font = _font(25)
    small_font = _font(20)
    centers: dict[str, tuple[int, int]] = {}
    boxes: dict[str, tuple[int, int, int, int]] = {}

    for lane_index, (lane_id, label, color) in enumerate(lane_entries):
        top = 56 + lane_index * 178
        bottom = top + 158
        draw.rounded_rectangle((22, top, width - 22, bottom), radius=8, fill="#FFFFFF", outline=color, width=3)
        draw.rectangle((22, top, 210, bottom), fill=_blend(color, 0.13), outline=color, width=2)
        draw.multiline_text((42, top + 53), _wrap(label, 7), font=font, fill="#26322B", spacing=4)
        nodes = grouped[lane_id]
        for node in nodes:
            x = 245 + node_order[node.id] * 230
            y = top + 39
            box = (x, y, x + 178, y + 80)
            boxes[node.id] = box
            centers[node.id] = ((box[0] + box[2]) // 2, (box[1] + box[3]) // 2)

    for edge in graph.edges:
        if edge.source not in centers or edge.target not in centers:
            continue
        _draw_arrow(draw, centers[edge.source], centers[edge.target], edge.label, small_font)

    for node in graph.nodes:
        box = boxes.get(node.id)
        if box is None:
            continue
        color = _node_color(node.type)
        if node.type == "decision":
            cx, cy = centers[node.id]
            polygon = [(cx, box[1] - 7), (box[2] + 8, cy), (cx, box[3] + 7), (box[0] - 8, cy)]
            draw.polygon(polygon, fill="#FFF9EB", outline=color, width=3)
        else:
            radius = 38 if node.type in {"start", "end"} else 5
            draw.rounded_rectangle(box, radius=radius, fill="#FFFFFF", outline=color, width=3)
        text = _wrap(node.label, 11)
        text_box = draw.multiline_textbbox((0, 0), text, font=small_font, align="center", spacing=3)
        text_width = text_box[2] - text_box[0]
        text_height = text_box[3] - text_box[1]
        cx, cy = centers[node.id]
        draw.multiline_text(
            (cx - text_width / 2, cy - text_height / 2 - 2),
            text,
            font=small_font,
            fill="#26322B",
            align="center",
            spacing=3,
        )
        if node.sap.tcodes:
            draw.text((box[0] + 7, box[3] - 18), node.sap.tcodes[0].code, font=_font(14), fill="#52796F")

    output = BytesIO()
    image.save(output, format="PNG", optimize=True)
    output.seek(0)
    return output


def _configure_document(document: DocumentType, model: BlueprintDocumentModel) -> None:
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    heading_tokens = {
        "Heading 1": (16, "2E74B5", 18, 10),
        "Heading 2": (13, "2E74B5", 14, 7),
        "Heading 3": (12, "1F4D78", 10, 5),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    header = section.header.paragraphs[0]
    header.text = f"SAP Blueprint | {model.project_name}"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _set_run_font(header.runs[0], size=8.5, color="68716B")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_run_font(footer.add_run(f"{model.process_name} | "), size=8.5, color="68716B")
    _append_page_field(footer)


def _add_title_block(document: DocumentType, model: BlueprintDocumentModel) -> None:
    kicker = document.add_paragraph("SAP BUSINESS BLUEPRINT")
    kicker.paragraph_format.space_after = Pt(4)
    _set_run_font(kicker.runs[0], size=9, color="2E74B5", bold=True)
    title = document.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(6)
    _set_run_font(title.add_run(model.graph.title), size=26, color="0B2545", bold=True)
    subtitle = document.add_paragraph(
        f"{model.graph.module} / {model.graph.process_scope} · {model.graph.sap_context.edition} {model.graph.sap_context.release}"
    )
    subtitle.paragraph_format.space_after = Pt(18)
    _set_run_font(subtitle.runs[0], size=12, color="68716B")
    table = document.add_table(rows=4, cols=2)
    _set_table_geometry(table, [2500, 6860])
    data = [
        ("项目", model.project_name),
        ("客户", model.customer_name or "未设置"),
        ("版本", f"发布 {model.release_no}" if model.release_no is not None else f"草稿修订 {model.revision_no}"),
        ("状态", model.lifecycle_state),
    ]
    for row, (label, value) in zip(table.rows, data, strict=True):
        _set_cell_text(row.cells[0], label, bold=True, color="1F4D78")
        _set_cell_text(row.cells[1], value)
    _mark_header_row(table.rows[0])
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def _add_project_section(document: DocumentType, model: BlueprintDocumentModel) -> None:
    _add_heading(document, "1. 项目与 SAP 环境", level=1)
    graph = model.graph
    table = document.add_table(rows=6, cols=2)
    _set_table_geometry(table, [2700, 6660])
    rows = [
        ("流程", model.process_name),
        ("模块 / 范围", f"{graph.module} / {graph.process_scope}"),
        ("SAP Edition / Release", f"{graph.sap_context.edition} / {graph.sap_context.release}"),
        ("部署方式", graph.sap_context.deployment),
        ("国家/地区", graph.sap_context.country),
        ("流程 ID", model.process_id),
    ]
    for row, (label, value) in zip(table.rows, rows, strict=True):
        _set_cell_text(row.cells[0], label, bold=True, color="1F4D78")
        _set_cell_text(row.cells[1], value)
    _mark_header_row(table.rows[0])


def _add_steps_section(document: DocumentType, graph: GraphDocument) -> None:
    _add_heading(document, "3. 流程步骤", level=1)
    table = document.add_table(rows=1, cols=6)
    _set_table_geometry(table, [520, 1300, 2200, 1200, 1500, 2640])
    _set_header_row(table.rows[0], ["#", "泳道", "步骤", "类型", "T-Code", "角色 / 状态"])
    lane_labels = {lane.id: lane.label for lane in graph.lanes}
    for index, node in enumerate(graph.nodes, 1):
        cells = table.add_row().cells
        values = [
            str(index),
            lane_labels.get(node.lane_id or "", "未分配"),
            node.label,
            str(node.sap.step_type or node.type),
            ", ".join(item.code for item in node.sap.tcodes) or "-",
            (", ".join(node.sap.roles) or "-") + f"\n{_metadata_summary(node)}",
        ]
        for cell, value in zip(cells, values, strict=True):
            _set_cell_text(cell, value, size=8.5)
    _set_table_geometry(table, [520, 1300, 2200, 1200, 1500, 2640])


def _add_evidence_section(document: DocumentType, graph: GraphDocument) -> None:
    _add_heading(document, "4. SAP 元数据与证据", level=1)
    rows = list(_evidence_rows(graph))
    if not rows:
        document.add_paragraph("当前修订尚无 SAP 元数据证据。")
        return
    table = document.add_table(rows=1, cols=4)
    _set_table_geometry(table, [2100, 2500, 1300, 3460])
    _set_header_row(table.rows[0], ["步骤", "对象", "状态", "证据引用"])
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row, strict=True):
            _set_cell_text(cell, value, size=8.5)
    _set_table_geometry(table, [2100, 2500, 1300, 3460])


def _add_gaps_section(document: DocumentType, graph: GraphDocument) -> None:
    _add_heading(document, "5. GAP List", level=1)
    gaps = [(node.label, node.sap.gap) for node in graph.nodes if node.sap.gap.status != "none"]
    if not gaps:
        document.add_paragraph("当前修订没有 GAP 候选或正式 GAP。")
        return
    table = document.add_table(rows=1, cols=5)
    _set_table_geometry(table, [1550, 1100, 2200, 2600, 1910])
    _set_header_row(table.rows[0], ["步骤", "状态", "差异描述", "建议方案", "负责人 / 证据"])
    for label, gap in gaps:
        cells = table.add_row().cells
        values = [
            label,
            _gap_status(gap.status),
            gap.description or "待补充",
            gap.recommendation or "待评估",
            f"{gap.owner or '未分配'}\n{', '.join(gap.evidence_refs) or '待确认'}",
        ]
        for cell, value in zip(cells, values, strict=True):
            _set_cell_text(cell, value, size=8.5)
    _set_table_geometry(table, [1550, 1100, 2200, 2600, 1910])


def _add_version_section(document: DocumentType, model: BlueprintDocumentModel) -> None:
    _add_heading(document, "6. 版本信息", level=1)
    table = document.add_table(rows=5, cols=2)
    _set_table_geometry(table, [2700, 6660])
    rows = [
        ("修订号", str(model.revision_no)),
        ("发布版本", str(model.release_no) if model.release_no is not None else "未发布"),
        ("生命周期", model.lifecycle_state),
        ("创建人", model.created_by),
        ("创建时间", model.created_at),
    ]
    for row, (label, value) in zip(table.rows, rows, strict=True):
        _set_cell_text(row.cells[0], label, bold=True, color="1F4D78")
        _set_cell_text(row.cells[1], value)
    _mark_header_row(table.rows[0])
    note = document.add_paragraph()
    note.paragraph_format.space_before = Pt(10)
    note.paragraph_format.space_after = Pt(0)
    _set_run_font(
        note.add_run(
            "注意：待确认字段和候选 GAP 必须由 SAP 顾问复核；本文档不会把模型建议自动升级为已验证结论。"
        ),
        size=9.5,
        color="7A5A00",
        bold=True,
    )


def _add_heading(document: DocumentType, text: str, *, level: int) -> None:
    document.add_paragraph(text, style=f"Heading {level}")


def _set_run_font(
    run, *, size: float, color: str = "252A27", bold: bool | None = None
) -> None:
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def _set_table_geometry(table, widths: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table_element = table._tbl
    properties = table_element.tblPr
    layout = properties.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        properties.append(layout)
    layout.set(qn("w:type"), "fixed")
    width = properties.find(qn("w:tblW"))
    if width is None:
        width = OxmlElement("w:tblW")
        properties.append(width)
    width.set(qn("w:type"), "dxa")
    width.set(qn("w:w"), str(sum(widths)))
    indent = properties.find(qn("w:tblInd"))
    if indent is None:
        indent = OxmlElement("w:tblInd")
        properties.append(indent)
    indent.set(qn("w:type"), "dxa")
    indent.set(qn("w:w"), "120")
    _set_table_borders(table)
    grid = table_element.tblGrid
    for child in list(grid):
        grid.remove(child)
    for value in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(value))
        grid.append(column)
    for row in table.rows:
        for cell, value in zip(row.cells, widths, strict=True):
            tc_width = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            tc_width.set(qn("w:type"), "dxa")
            tc_width.set(qn("w:w"), str(value))
            _set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _set_table_borders(table) -> None:
    properties = table._tbl.tblPr
    borders = properties.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    else:
        for child in list(borders):
            borders.remove(child)
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement(f"w:{name}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:color"), "C8D0D6")
        borders.append(edge)


def _set_cell_margins(cell) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    else:
        for child in list(margins):
            margins.remove(child)
    for name, value in (("top", 80), ("bottom", 80), ("start", 120), ("end", 120)):
        node = OxmlElement(f"w:{name}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)


def _set_header_row(row, values: list[str]) -> None:
    _mark_header_row(row)
    for cell, value in zip(row.cells, values, strict=True):
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "E8EEF5")
        cell._tc.get_or_add_tcPr().append(shading)
        _set_cell_text(cell, value, bold=True, color="1F4D78", size=8.5)


def _mark_header_row(row) -> None:
    row_properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    row_properties.append(repeat)


def _set_cell_text(
    cell, text: str, *, bold: bool = False, color: str = "252A27", size: float = 9.5
) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.clear()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.15
    _set_run_font(paragraph.add_run(str(text)), size=size, color=color, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _append_page_field(paragraph) -> None:
    run = paragraph.add_run("Page ")
    _set_run_font(run, size=8.5, color="68716B")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    field_run = paragraph.add_run()
    field_run._r.extend([begin, instruction, separate, text, end])
    _set_run_font(field_run, size=8.5, color="68716B")


def _mermaid_lines(graph: GraphDocument) -> list[str]:
    lane_nodes: dict[str, list[tuple[int, object]]] = {}
    lane_labels = {lane.id: lane.label for lane in graph.lanes}
    for index, node in enumerate(graph.nodes):
        lane_nodes.setdefault(node.lane_id or "__unassigned__", []).append((index, node))
    lines = ["```mermaid", "flowchart LR" if graph.direction == "LR" else "flowchart TB"]
    for lane_index, (lane_id, nodes) in enumerate(lane_nodes.items()):
        label = lane_labels.get(lane_id, "未分配")
        lines.append(f'  subgraph L{lane_index}["{_mermaid(label)}"]')
        for index, node in nodes:
            lines.append(f"    N{index}{_mermaid_shape(node.type, node.label)}")
        lines.append("  end")
    index_by_id = {node.id: index for index, node in enumerate(graph.nodes)}
    for edge in graph.edges:
        if edge.source not in index_by_id or edge.target not in index_by_id:
            continue
        label = f'|"{_mermaid(edge.label)}"|' if edge.label else ""
        lines.append(f"  N{index_by_id[edge.source]} -->{label} N{index_by_id[edge.target]}")
    lines.append("```")
    return lines


def _mermaid_shape(node_type: object, label: str) -> str:
    safe = _mermaid(label)
    if str(node_type) in {"start", "end"}:
        return f'(["{safe}"])'
    if str(node_type) == "decision":
        return f'{{"{safe}"}}'
    if str(node_type) == "subprocess":
        return f'[["{safe}"]]'
    return f'["{safe}"]'


def _evidence_rows(graph: GraphDocument) -> Iterable[tuple[str, str, str, str]]:
    for node in graph.nodes:
        for item in node.sap.tcodes:
            yield node.label, f"T-Code {item.code}", _metadata_status(item.status), item.evidence_ref or "待确认"
        for item in node.sap.fiori_apps:
            yield node.label, f"Fiori {item.name}", _metadata_status(item.status), item.evidence_ref or "待确认"
        for item in node.sap.configuration_points:
            yield node.label, item.label, _metadata_status(item.status), item.evidence_ref or "待确认"
        for item in node.sap.best_practice_refs:
            yield node.label, f"{item.scope_item} / {item.step}", "引用", item.evidence_ref or "待确认"


def _metadata_summary(node) -> str:
    statuses = [item.status for item in node.sap.tcodes]
    statuses.extend(item.status for item in node.sap.fiori_apps)
    statuses.extend(item.status for item in node.sap.configuration_points)
    if not statuses:
        return "未设置"
    if any(str(status) == "pending_confirmation" for status in statuses):
        return "待确认"
    if any(str(status) == "suggested" for status in statuses):
        return "建议"
    return "已验证"


def _metadata_status(status: object) -> str:
    return {
        "suggested": "建议",
        "pending_confirmation": "待确认",
        "verified": "已验证",
    }.get(str(status), str(status))


def _gap_status(status: object) -> str:
    return {
        "candidate": "候选",
        "confirmed": "已确认",
        "rejected": "已驳回",
        "resolved": "已解决",
        "none": "无 GAP",
    }.get(str(status), str(status))


def _md(value: object) -> str:
    return str(value).replace("|", "\\|").replace("\n", "<br>")


def _mermaid(value: object) -> str:
    return str(value or "").replace('"', "'").replace("\n", " ")


def _font(size: int):
    candidates = [
        Path("C:/Windows/Fonts/msyh.ttc"),
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


def _wrap(value: str, length: int) -> str:
    return "\n".join(value[index : index + length] for index in range(0, len(value), length))


def _node_color(node_type: object) -> str:
    return {
        "start": "#2D7A5E",
        "end": "#AD4C4C",
        "task": "#3E6F98",
        "decision": "#C58B2C",
        "subprocess": "#745D91",
    }.get(str(node_type), "#68736D")


def _blend(color: str, ratio: float) -> str:
    value = color.lstrip("#")
    red, green, blue = (int(value[index : index + 2], 16) for index in (0, 2, 4))
    red = round(255 - (255 - red) * ratio)
    green = round(255 - (255 - green) * ratio)
    blue = round(255 - (255 - blue) * ratio)
    return f"#{red:02X}{green:02X}{blue:02X}"


def _draw_arrow(draw, source, target, label: str | None, font) -> None:
    sx, sy = source
    tx, ty = target
    start = (sx + (48 if tx >= sx else -48), sy)
    end = (tx - (48 if tx >= sx else -48), ty)
    draw.line((start, end), fill="#68736D", width=3)
    direction = 1 if end[0] >= start[0] else -1
    arrow = [(end[0], end[1]), (end[0] - direction * 12, end[1] - 7), (end[0] - direction * 12, end[1] + 7)]
    draw.polygon(arrow, fill="#68736D")
    if label:
        midpoint = ((start[0] + end[0]) // 2, (start[1] + end[1]) // 2 - 22)
        draw.text(midpoint, label, font=font, fill="#4F5852", anchor="mm")
